"""Export berlapis (docs/INTENT-OUTPUT.md §8.3): XLSX = data mentah + sumber; DOCX = memo + tabel; PDF = HTML print."""
from __future__ import annotations

import io
import json
from typing import Any


def _panel_order(result: dict[str, Any]) -> list[dict[str, Any]]:
    return result.get("panels") or []


def build_xlsx(result: dict[str, Any]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="1B4DD8")

    def sheet(name: str):
        ws = wb.create_sheet(name[:31]) if wb.sheetnames != ["Sheet"] else wb.active
        if ws.title == "Sheet":
            ws.title = name[:31]
        return ws

    def write_table(ws, start_row: int, table: dict) -> int:
        columns = table.get("columns") or []
        rows = table.get("rows") or []
        for j, col in enumerate(columns, start=1):
            cell = ws.cell(row=start_row, column=j, value=col.get("label", col.get("key")))
            cell.font, cell.fill = head_font, head_fill
            cell.alignment = Alignment(horizontal="center")
        for i, row in enumerate(rows, start=start_row + 1):
            for j, col in enumerate(columns, start=1):
                ws.cell(row=i, column=j, value=row.get(col["key"]))
        return start_row + 1 + len(rows) + 1

    l0 = result.get("l0") or {}
    memo = result.get("memo") or {}
    ws = sheet("Ringkasan")
    ws.append([l0.get("title", "IDXMACA")]); ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Query: {l0.get('query', '')}"])
    ws.append([f"Persona: {l0.get('persona', '')}"])
    ws.append([])
    ws.append(["KPI", "Nilai", "Konteks"])
    for kpi in l0.get("kpis", []):
        ws.append([kpi.get("label"), kpi.get("value"), kpi.get("sub")])
    ws.append([])
    ws.append(["Poin kunci"])
    for bullet in l0.get("bullets", []):
        ws.append([bullet])
    ws.append([])
    ws.append(["Memo 3 lensa"])
    for lens in memo.get("lenses", []):
        ws.append([lens.get("title"), lens.get("text")])
    ws.append([])
    ws.append([memo.get("disclaimer", "")])
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 60
    ws.column_dimensions["C"].width = 40

    for panel in _panel_order(result):
        ws = sheet(panel["id"] + " " + panel["title"])
        ws.append([panel["id"] + " · " + panel["title"]]); ws["A1"].font = Font(bold=True, size=13)
        for piece in panel.get("items", []):
            ws.append([])
            ws.append([piece.get("io"), piece.get("title")]); ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
            ws.append([piece.get("narrative", "")])
            if piece.get("table"):
                write_table(ws, ws.max_row + 1, piece["table"])
            for source in (piece.get("evidence") or [])[:1]:
                ws.append([f"Bukti: {source}"])
        ws.column_dimensions["A"].width = 22
        for col in range(2, 9):
            ws.column_dimensions[get_column_letter(col)].width = 22

    ws = sheet("Komoditas-Gabungan") if False else None  # pragma: no cover
    ws = sheet("Evidence")
    ws.append(["Evidence ID", "Label", "Endpoint", "Params", "Sumber", "Ditarik", "Nilai"])
    for cell in ws[1]:
        cell.font, cell.fill = head_font, head_fill
    for ev_id, ev in (result.get("evidence") or {}).items():
        ws.append([ev_id, ev.get("label"), ev.get("endpoint"), json.dumps(ev.get("params", {})),
                   ev.get("source"), ev.get("fetched_at"), json.dumps(ev.get("value"), default=str)[:500]])
    for col, width in zip("ABCDEFG", (12, 40, 34, 40, 14, 22, 60)):
        ws.column_dimensions[col].width = width

    if "Sheet" in wb.sheetnames and wb["Sheet"].max_row == 1 and wb["Sheet"].max_column == 1:
        del wb["Sheet"]
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def build_docx(result: dict[str, Any]) -> bytes:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    l0 = result.get("l0") or {}
    memo = result.get("memo") or {}
    doc.add_heading("IDXMACA — Memo Analisis", level=0)
    doc.add_paragraph(f"Query: {l0.get('query', '')}")
    doc.add_paragraph(f"Persona: {l0.get('persona', '')} · Playbook: {', '.join(l0.get('playbooks', []))}")
    doc.add_heading("Ringkasan Eksekutif", level=1)
    doc.add_paragraph(l0.get("title", ""))
    for bullet in l0.get("bullets", []):
        doc.add_paragraph(bullet, style="List Bullet")
    for lens in memo.get("lenses", []):
        doc.add_heading(lens.get("title", "Lensa"), level=2)
        doc.add_paragraph(lens.get("text", ""))
    for panel in _panel_order(result):
        doc.add_heading(f"{panel['id']} · {panel['title']}", level=1)
        for piece in panel.get("items", []):
            if piece.get("status") != "ready":
                continue
            doc.add_heading(piece.get("title", ""), level=3)
            if piece.get("narrative"):
                doc.add_paragraph(piece["narrative"])
            table = piece.get("table")
            if table and table.get("rows"):
                cols = table.get("columns") or []
                table_obj = doc.add_table(rows=1, cols=len(cols))
                table_obj.style = "Light Grid Accent 1"
                for i, col in enumerate(cols):
                    table_obj.rows[0].cells[i].text = str(col.get("label", col.get("key")))
                for row in table["rows"][:30]:
                    cells = table_obj.add_row().cells
                    for i, col in enumerate(cols):
                        cells[i].text = str(row.get(col["key"], "—"))
                if table.get("note"):
                    doc.add_paragraph(table["note"], style="Intense Quote")
    doc.add_paragraph(memo.get("disclaimer", ""))
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def build_print_html(result: dict[str, Any]) -> str:
    l0 = result.get("l0") or {}
    memo = result.get("memo") or {}
    parts = [f"""<!DOCTYPE html><html lang="id"><head><meta charset="utf-8"><title>IDXMACA — {l0.get('title', 'Memo')}</title>
<style>
body{{font-family:Roboto,Arial,sans-serif;color:#1f1f1f;margin:32px auto;max-width:900px;line-height:1.5}}
h1{{font-size:24px}} h2{{font-size:18px;margin-top:28px}} h3{{font-size:15px;margin-bottom:4px}}
.kpis{{display:flex;gap:12px;flex-wrap:wrap}}.kpi{{border:1px solid #e3e3e3;border-radius:10px;padding:10px 14px;min-width:140px}}
.kpi b{{display:block;font-size:18px}} .src{{font-size:11px;color:#747775}}
table{{border-collapse:collapse;width:100%;font-size:12px;margin:8px 0}}td,th{{border:1px solid #e3e3e3;padding:6px 8px;text-align:right}}
th:first-child,td:first-child{{text-align:left}} .disc{{background:#f0f4f9;border-radius:10px;padding:12px;font-size:12px;margin-top:24px}}
</style></head><body>
<h1>IDXMACA — Memo Analisis</h1><p>{l0.get('query','')}<br><span class="src">{l0.get('persona','')} · {', '.join(l0.get('playbooks', []))}</span></p>
<h2>Ringkasan</h2><p><b>{l0.get('title','')}</b></p><div class="kpis">"""]
    for kpi in l0.get("kpis", []):
        parts.append(f"<div class='kpi'><span class='src'>{kpi.get('label','')}</span><b>{kpi.get('value','')}</b><span class='src'>{kpi.get('sub','') or ''}</span></div>")
    parts.append("</div><ul>")
    for bullet in l0.get("bullets", []):
        parts.append(f"<li>{bullet}</li>")
    parts.append("</ul>")
    for lens in memo.get("lenses", []):
        parts.append(f"<h2>{lens.get('title','')}</h2><p>{lens.get('text','')}</p>")
    for panel in _panel_order(result):
        parts.append(f"<h2>{panel['id']} · {panel['title']}</h2>")
        for piece in panel.get("items", []):
            if piece.get("status") != "ready":
                continue
            parts.append(f"<h3>{piece.get('title','')}</h3><p>{piece.get('narrative','')}</p>")
            table = piece.get("table")
            if table and table.get("rows"):
                parts.append("<table><tr>" + "".join(f"<th>{c.get('label', c.get('key'))}</th>" for c in table["columns"]) + "</tr>")
                for row in table["rows"][:40]:
                    parts.append("<tr>" + "".join(f"<td>{row.get(c['key'], '—') if row.get(c['key']) is not None else '—'}</td>" for c in table["columns"]) + "</tr>")
                parts.append("</table>")
            if table and table.get("note"):
                parts.append(f"<p class='src'>{table['note']}</p>")
    parts.append(f"<div class='disc'>{memo.get('disclaimer','')}</div></body></html>")
    return "".join(parts)